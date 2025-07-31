import os
import re
import logging
import pdfplumber
import pandas as pd
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from .text_processing import clean_text
from .config import RAGConfig

logger = logging.getLogger(__name__)

def extract_pdf(file_path, config: RAGConfig):
    all_content = []
    base_filename = os.path.basename(file_path)
    try:
        # Try to open as a real PDF
        with pdfplumber.open(file_path) as pdf:
            total_pages = len(pdf.pages)
            logger.info(f"Extracting text/tables from PDF: {base_filename} ({total_pages} pages) - Applying structural heuristics.")

            toc_pattern = re.compile(r"(\.|\s){4,}\s*\d+\s*$")
            intro_keywords = re.compile(r'\b(introduction|about sat ?sure|company profile|executive summary|overview|problem statement|client need|challenge)\b', re.IGNORECASE)
            closing_keywords = re.compile(r'\b(appendix|annexure|pricing|cost|timeline|schedule|next steps|conclusion)\b', re.IGNORECASE)

            for page_num, page in enumerate(pdf.pages):
                current_page = page_num + 1
                text = page.extract_text(layout="normal") or ""
                cleaned_text = clean_text(text)
                page_word_count = len(cleaned_text.split())

                potential_section = "main_content"
                is_first_page = (page_num == 0)
                is_likely_toc_page = (page_num == 1 or page_num == 2)
                is_near_end = (page_num >= total_pages - 3)

                if is_first_page:
                    potential_section = "title_page"
                    potential_title = ""
                    lines = [line.strip() for line in text.split('\n') if line.strip()]
                    if lines:
                        first_few_lines = lines[:7]
                        if first_few_lines:
                            potential_title = max(first_few_lines, key=len)
                            if len(potential_title) > 80 or potential_title.islower():
                                potential_title = first_few_lines[0]
                    logger.debug(f"PDF Pg {current_page}: Marked as 'title_page'. Potential Title: '{potential_title[:50]}...'")

                elif is_likely_toc_page:
                    lines = text.split('\n')
                    toc_lines_count = sum(1 for line in lines if toc_pattern.search(line.strip()))
                    if toc_lines_count >= 3 and page_word_count < 300 :
                        potential_section = "table_of_contents"
                        logger.debug(f"PDF Pg {current_page}: Marked as 'table_of_contents' (Found {toc_lines_count} matching lines).")

                if potential_section == "main_content" and page_num < 5:
                    if intro_keywords.search(cleaned_text):
                        potential_section = "introduction_overview"
                        logger.debug(f"PDF Pg {current_page}: Marked as 'introduction_overview' based on keywords.")
                    elif page_num < 3:
                         potential_section = "front_matter"

                if potential_section == "main_content" and is_near_end:
                     if closing_keywords.search(cleaned_text):
                         potential_section = "closing_appendix"
                         logger.debug(f"PDF Pg {current_page}: Marked as 'closing_appendix' based on keywords.")
                     elif page_num == total_pages -1:
                         potential_section = "final_page"

                source_info = {"page": current_page}
                if is_first_page and potential_title:
                     source_info["potential_doc_title"] = potential_title

                if cleaned_text:
                    all_content.append({
                        "type": "text", "content": cleaned_text,
                        "source_info": source_info, "file_type": "pdf",
                        "potential_section": potential_section
                    })

                try:
                    for table_idx, table_content in enumerate(page.extract_tables()):
                         if table_content:
                             table_df = pd.DataFrame(table_content)
                             if not table_df.empty and not pd.api.types.is_numeric_dtype(table_df.iloc[0].dropna()):
                                 try:
                                     table_df.columns = table_df.iloc[0].fillna('').astype(str)
                                     table_df = table_df[1:]
                                 except Exception: pass
                             table_df = table_df.fillna('')
                             table_string = table_df.to_string(index=False, header=True)
                             cleaned_table_string = clean_text(table_string)
                             if cleaned_table_string:
                                 all_content.append({
                                    "type": "table", "content": cleaned_table_string,
                                    "source_info": {**source_info, "table_index_on_page": table_idx + 1},
                                    "file_type": "pdf", "potential_section": potential_section
                                 })
                except Exception as e_table:
                    logger.warning(f"Could not extract/process table on page {current_page} in {base_filename}: {e_table}")

        logger.info(f"Extracted {len(all_content)} text/table content blocks from PDF: {base_filename} with structural tagging.")
        return all_content
    except Exception:
        logger.warning(f"Could not open {base_filename} as a PDF. Attempting to read as plain text.")
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                text = f.read()
            cleaned_text = clean_text(text)
            if cleaned_text:
                all_content.append({
                    "type": "text", "content": cleaned_text,
                    "source_info": {"page": 1}, "file_type": "pdf", # Pretend it's a PDF
                    "potential_section": "main_content"
                })
            return all_content
        except Exception as e_text:
            logger.error(f"Failed to read {base_filename} as plain text: {e_text}", exc_info=True)
            return []

def extract_pptx(file_path, config: RAGConfig):
    all_content = []
    base_filename = os.path.basename(file_path)
    max_chars = config.max_chars_per_element
    merge_threshold_words = config.pptx_merge_threshold_words
    try:
        # Try to open as a real PPTX
        prs = Presentation(file_path)
        logger.info(f"Extracting text from PPTX: {base_filename} ({len(prs.slides)} slides)")
        pending_title_slide_data = None
        for i, slide in enumerate(prs.slides):
             current_slide_number = i + 1
             current_slide_title_text = ""
             current_slide_other_texts = []
             current_slide_has_title_placeholder = False
             title_shape = None
             try:
                if slide.shapes.title:
                    title_shape = slide.shapes.title
                    current_slide_has_title_placeholder = True
             except AttributeError: pass

             if title_shape and title_shape.has_text_frame:
                 cleaned_title = clean_text(title_shape.text_frame.text)
                 current_slide_title_text = cleaned_title if cleaned_title else ""

             for shape in slide.shapes:
                 if shape == title_shape: continue
                 is_placeholder = shape.is_placeholder
                 is_body_placeholder = False
                 if is_placeholder:
                     try:
                         ph_type = shape.placeholder_format.type
                         is_body_placeholder = ph_type in [MSO_SHAPE_TYPE.BODY, MSO_SHAPE_TYPE.OBJECT, MSO_SHAPE_TYPE.SUBTITLE, MSO_SHAPE_TYPE.CONTENT, MSO_SHAPE_TYPE.TEXT_BOX, MSO_SHAPE_TYPE.CHART, MSO_SHAPE_TYPE.TABLE, MSO_SHAPE_TYPE.PICTURE]
                     except AttributeError: pass
                 if shape.has_text_frame:
                     text = shape.text_frame.text
                     cleaned = clean_text(text)
                     if cleaned:
                          if len(cleaned) > max_chars:
                              cleaned = cleaned[:max_chars] + "...(truncated shape)"
                          prefix = "[Body]: " if is_body_placeholder and not current_slide_title_text and not current_slide_other_texts else ""
                          current_slide_other_texts.append(prefix + cleaned)

             if slide.has_notes_slide:
                try:
                     notes_text = slide.notes_slide.notes_text_frame.text
                     cleaned_notes = clean_text(notes_text)
                     if cleaned_notes:
                         if len(cleaned_notes) > max_chars * 2:
                             cleaned_notes = cleaned_notes[:max_chars*2] + "...(truncated notes)"
                         current_slide_other_texts.append(f"[Notes]: {cleaned_notes}")
                except Exception as e_notes:
                    logger.warning(f"Could not extract notes from slide {current_slide_number}: {e_notes}")

             current_slide_full_content = current_slide_title_text
             if current_slide_other_texts:
                 current_slide_full_content += ("\n" if current_slide_full_content else "") + "\n".join(current_slide_other_texts)
             current_slide_full_content = current_slide_full_content.strip()
             other_text_word_count = sum(len(s.split()) for s in current_slide_other_texts)

             merged_content_block = None
             should_merge = False
             if pending_title_slide_data:
                 if pending_title_slide_data['has_title'] and pending_title_slide_data['other_words'] <= merge_threshold_words and current_slide_full_content:
                     should_merge = True
                     logger.info(f"Merging slide {pending_title_slide_data['number']} (title-like) with content from slide {current_slide_number}.")

             if should_merge:
                merged_text = f"[Title from Slide {pending_title_slide_data['number']}]: {pending_title_slide_data['title']}\n\n[Content from Slide {pending_title_slide_data['number']} (if any)]:\n{pending_title_slide_data['content_without_title']}\n\n---\n\n[Content from Slide {current_slide_number}]:\n{current_slide_full_content}"
                merged_content_block = {"type": "slide_text_merged", "content": merged_text.strip(), "source_info": {"slide_title": pending_title_slide_data['number'], "slide_content": current_slide_number}, "file_type": "pptx"}
                all_content.append(merged_content_block)
                pending_title_slide_data = None
             else:
                if pending_title_slide_data:
                     if pending_title_slide_data['content']:
                         all_content.append({"type": "slide_text", "content": pending_title_slide_data['content'], "source_info": {"slide": pending_title_slide_data['number']}, "file_type": "pptx"})
                     pending_title_slide_data = None

                if current_slide_has_title_placeholder and current_slide_title_text and other_text_word_count <= merge_threshold_words and current_slide_full_content:
                    logger.debug(f"Slide {current_slide_number} ('{current_slide_title_text[:30]}...') is potential title slide for next content. Holding.")
                    content_without_title_parts = [txt for txt in current_slide_other_texts if not txt.startswith(current_slide_title_text)]
                    content_w_o_title_str = "\n".join(content_without_title_parts).strip()
                    pending_title_slide_data = {
                        "content": current_slide_full_content, "content_without_title": content_w_o_title_str,
                        "number": current_slide_number, "has_title": current_slide_has_title_placeholder,
                        "other_words": other_text_word_count, "title": current_slide_title_text
                    }
                else:
                    if current_slide_full_content:
                        all_content.append({"type": "slide_text", "content": current_slide_full_content, "source_info": {"slide": current_slide_number}, "file_type": "pptx"})

        if pending_title_slide_data:
             logger.debug(f"Processing final pending title slide {pending_title_slide_data['number']} at end.")
             if pending_title_slide_data['content']:
                 all_content.append({"type": "slide_text", "content": pending_title_slide_data['content'], "source_info": {"slide": pending_title_slide_data['number']}, "file_type": "pptx"})

        logger.info(f"Extracted {len(all_content)} content blocks from PPTX {base_filename} (Merge strategy applied).")
        return all_content
    except Exception:
        logger.warning(f"Could not open {base_filename} as a PPTX. Attempting to read as plain text.")
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                text = f.read()
            cleaned_text = clean_text(text)
            if cleaned_text:
                all_content.append({
                    "type": "text", "content": cleaned_text,
                    "source_info": {"slide": 1}, "file_type": "pptx", # Pretend it's a PPTX
                    "potential_section": "main_content"
                })
            return all_content
        except Exception as e_text:
            logger.error(f"Failed to read {base_filename} as plain text: {e_text}", exc_info=True)
            return []
